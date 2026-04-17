"""Parser for sales report Excel from 1C.
Structure: hierarchical rows with indent levels 0,2,4,6,8,10.
Columns per branch: qty (col+0), amount (col+1), price (col+2), sum_no_nds (col+3).
"""
import re
from pathlib import Path
from datetime import date

import openpyxl

BRANCH_COLS = [
    ("BEREКЕ",     3,  4),   # ИП Береке (col 3 = qty, col 4 = sum)
    ("MAIN",       7,  8),   # Основной склад
    ("AKTAU",     11, 12),
    ("AKTOBE",    15, 16),
    ("ALMATY",    19, 20),
    ("ASTANA",    23, 24),
    ("ATYRAU",    27, 28),
    ("KARAGANDA", 31, 32),
    ("KOKSHETAU", 35, 36),
    ("SEMEY",     39, 40),
    ("SHYMKENT",  43, 44),
]

# Section names to skip entirely (header + all descendants)
SKIP_SECTIONS = re.compile(
    r"^(не использовать|материалы|на удаление|нет в продаже|"
    r"услуги сторонних организаций|услуги|итого)\s*$",
    re.I,
)
# Codes that trigger section skip
SKIP_CODES = {"KSN00000263", "URA00000435", "AST00000683", "00000000005", "00000000008"}

# Skip just this row (header only, keep children)
SKIP_HEADER_ONLY = re.compile(r"^(товары)\s*$", re.I)
SKIP_HEADER_CODES = {"00000000002"}

BONUS_CODE = "00000000012"

# Regex to detect and strip "БОНУСЫ / Бонусы / бонусы" suffix (case-insensitive)
_BONUS_SUFFIX_RE = re.compile(r'\s+бонус[ыь]?\s*$', re.IGNORECASE)

# Cat1 names to normalize (lowercase key → canonical value)
_CAT1_NORMALIZE: dict[str, str] = {
    "кухмастер": "КУХМАСТЕР",
}

RU_MONTHS = {
    "январ": 1, "феврал": 2, "март": 3, "апрел": 4,
    "май": 5, "мая": 5, "июн": 6, "июл": 7, "август": 8,
    "сентябр": 9, "октябр": 10, "ноябр": 11, "декабр": 12,
}


def _extract_period_date(filename: str) -> date:
    name = Path(filename).stem.lower()
    # Try "март 31" or "31 март" style
    for ru, num in RU_MONTHS.items():
        if ru in name:
            day_match = re.search(r"\b(\d{1,2})\b", name)
            day = int(day_match.group(1)) if day_match else 1
            year = date.today().year
            return date(year, num, day)
    # Try dd.mm.yy
    m = re.search(r"(\d{1,2})[.\-](\d{1,2})[.\-](\d{2,4})", name)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if y < 100:
            y += 2000
        return date(y, mo, d)
    return date.today()


def parse_sales_report(filepath: str) -> list[dict]:
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.sheet_by_name = None
    ws = wb.active

    period_date = _extract_period_date(Path(filepath).name)

    # Read all raw rows
    raw = []
    for r in range(3, ws.max_row + 1):
        cell = ws.cell(r, 2)
        name = str(cell.value).strip() if cell.value else ""
        code = str(ws.cell(r, 1).value).strip() if ws.cell(r, 1).value else ""
        indent = float(cell.alignment.indent) if cell.alignment else 0.0

        branch_data = {}
        for bc, qc, sc in BRANCH_COLS:
            qty = ws.cell(r, qc).value
            amt = ws.cell(r, sc).value
            if qty or amt:
                branch_data[bc] = (float(qty or 0), float(amt or 0))

        raw.append({
            "r": r,
            "code": code,
            "name": name,
            "indent": indent,
            "branch_data": branch_data,
        })

    # Mark leaf nodes: leaf if next row has indent <= current
    for i in range(len(raw)):
        next_indent = raw[i + 1]["indent"] if i + 1 < len(raw) else -1
        raw[i]["is_leaf"] = next_indent <= raw[i]["indent"]

    # Traverse and build entries
    entries = []
    # category context: indent_level -> name
    ctx: dict[float, str] = {}
    # skip_until_indent: if set, skip all rows with indent > this value
    skip_until: float | None = None
    in_bonus = False
    bonus_indent: float | None = None  # indent level of the БОНУСЫ header row

    for row in raw:
        indent = row["indent"]
        name = row["name"]
        code = row["code"]

        # Clear stale context levels
        for lvl in list(ctx.keys()):
            if lvl >= indent:
                del ctx[lvl]

        # Handle skip_until
        if skip_until is not None:
            if indent > skip_until:
                continue  # still in skipped section
            else:
                skip_until = None  # exited skip section

        # Check if this row starts a skip section
        if SKIP_SECTIONS.match(name) or code in SKIP_CODES:
            skip_until = indent
            continue

        # Skip ТОВАРЫ header only
        if SKIP_HEADER_ONLY.match(name) or code in SKIP_HEADER_CODES:
            continue

        # Track bonus section: enter on BONUS_CODE row, exit when back at same indent
        if code == BONUS_CODE:
            in_bonus = True
            bonus_indent = indent
        elif in_bonus and bonus_indent is not None and indent <= bonus_indent:
            in_bonus = False
            bonus_indent = None

        # Update context
        ctx[indent] = name

        # Only save leaf nodes
        if not row["is_leaf"]:
            continue

        # Skip rows with empty name (summary/header rows)
        if not name:
            continue

        # Build hierarchy (from context at lower indent levels)
        sorted_levels = sorted(k for k in ctx.keys() if k < indent)
        if in_bonus:
            # Skip the БОНУСЫ level (indent=2) — use indent=4+ as categories
            cats = [ctx[k] for k in sorted_levels if k >= 4]
        else:
            cats = [ctx[k] for k in sorted_levels if k >= 2]  # skip indent=0 (ТОВАРЫ)

        # Strip "БОНУСЫ / Бонусы" suffix (case-insensitive) from category names
        def _clean(s: str | None) -> str | None:
            if s is None:
                return None
            return _BONUS_SUFFIX_RE.sub('', s).strip() or s

        # Normalize cat1 case variants (e.g. "Кухмастер" → "КУХМАСТЕР")
        def _normalize(s: str | None) -> str | None:
            if s is None:
                return None
            return _CAT1_NORMALIZE.get(s.strip().lower(), s)

        # Detect bonus-by-category: top-level category ending with "БОНУСЫ"
        orig_cat1 = cats[0] if cats else None
        is_cat_bonus = bool(orig_cat1 and _BONUS_SUFFIX_RE.search(orig_cat1))

        cat1 = _normalize(_clean(cats[0])) if len(cats) > 0 else None
        cat2 = _clean(cats[1]) if len(cats) > 1 else None
        cat3 = _clean(cats[2]) if len(cats) > 2 else None
        cat4 = _clean(cats[3]) if len(cats) > 3 else None

        for bc, (qty, amt) in row["branch_data"].items():
            if qty == 0 and amt == 0:
                continue
            entries.append({
                "period_date": period_date,
                "code": code,
                "name": name,
                "cat1": cat1,
                "cat2": cat2,
                "cat3": cat3,
                "cat4": cat4,
                "is_bonus": in_bonus or is_cat_bonus,
                "branch_code": bc,
                "qty": qty,
                "amount": amt,
            })

    return entries
